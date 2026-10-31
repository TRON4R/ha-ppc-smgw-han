"""Tests for dated Saturday/Sunday/holiday zone rules."""

from __future__ import annotations

from datetime import date, datetime, time

import pytest
from freezegun import freeze_time
from homeassistant.core import HomeAssistant
from openpyxl import load_workbook
from pytest_homeassistant_custom_component.common import MockConfigEntry

from custom_components.smgw_han.aggregation import build_daily_summary
from custom_components.smgw_han import day_rules, zone_schedule
from custom_components.smgw_han.config_flow import (
    ZoneDefinitionError,
    _parse_optional_tariff_zones,
)
from custom_components.smgw_han.const import (
    CONF_HOLIDAYS,
    CONF_METER_ID,
    CONF_TARIFF_ZONES,
    CONF_TARIFF_ZONES_SATURDAY,
    CONF_TARIFF_ZONES_SUNDAY,
    CONF_ZONE_SCHEDULE,
    DOMAIN,
    OBIS_IMPORT,
    STORE_VERSION,
    TARIFF_TEMPLATE_WEEKEND,
    TARIFF_TEMPLATES_DAY_TYPES,
)
from custom_components.smgw_han.coordinator import SmgwCoordinator
from custom_components.smgw_han.export_files import write_xlsx
from custom_components.smgw_han.sensor import _dynamic_descriptions
from custom_components.smgw_han.smgw_client import (
    DailyData,
    MeterReading,
    SmgwConnectionError,
)


@pytest.fixture(autouse=True)
def _enable(enable_custom_integrations):
    yield


BW = TARIFF_TEMPLATES_DAY_TYPES[TARIFF_TEMPLATE_WEEKEND]
WEEKDAY = BW[CONF_TARIFF_ZONES]  # NT 00-06, HT 06-22, NT 22-24
SATURDAY = BW[CONF_TARIFF_ZONES_SATURDAY]  # NT 00-06, HT 06-13, NT 13-24
SUNDAY = BW[CONF_TARIFF_ZONES_SUNDAY]  # NT all day

# 2026-10-05 Mon, 10-10 Sat, 10-11 Sun; 10-03 is a Saturday AND a holiday
MONDAY_D, SATURDAY_D, SUNDAY_D = date(2026, 10, 5), date(2026, 10, 10), date(2026, 10, 11)
HOLIDAY_ON_SATURDAY = date(2026, 10, 3)
HOLIDAY_ON_WEEKDAY = date(2026, 12, 25)  # Friday


distinct_zone_names = day_rules.distinct_zone_names


def is_national_holiday(day: date) -> bool:
    return day_rules.is_holiday("national", day)


def day_zones_config(data: dict, _fallback, day: date) -> list[dict]:
    """Zones of ``day`` for flat entry data (test shorthand)."""
    return day_rules.zones_for_day(day_rules.entry_layout(data), day)


def _data(holidays: bool | str = True, **extra) -> dict:
    if holidays is True:
        holidays = "national"
    elif holidays is False:
        holidays = "off"
    return {
        CONF_TARIFF_ZONES: WEEKDAY,
        CONF_TARIFF_ZONES_SATURDAY: SATURDAY,
        CONF_TARIFF_ZONES_SUNDAY: SUNDAY,
        CONF_HOLIDAYS: holidays,
        **extra,
    }


# --- day-type resolution -------------------------------------------------


def test_weekday_uses_base_layout():
    assert day_zones_config(_data(), WEEKDAY, MONDAY_D) == WEEKDAY


def test_saturday_and_sunday_use_their_override():
    assert day_zones_config(_data(), WEEKDAY, SATURDAY_D) == SATURDAY
    assert day_zones_config(_data(), WEEKDAY, SUNDAY_D) == SUNDAY


def test_empty_override_falls_back_to_base():
    data = {CONF_TARIFF_ZONES: WEEKDAY}
    assert day_zones_config(data, WEEKDAY, SATURDAY_D) == WEEKDAY
    assert day_zones_config(data, WEEKDAY, SUNDAY_D) == WEEKDAY


def test_holiday_on_weekday_is_sunday_when_enabled():
    assert is_national_holiday(HOLIDAY_ON_WEEKDAY)
    assert day_zones_config(_data(True), WEEKDAY, HOLIDAY_ON_WEEKDAY) == SUNDAY


def test_holiday_ignored_when_switch_is_off():
    assert day_zones_config(_data(False), WEEKDAY, HOLIDAY_ON_WEEKDAY) == WEEKDAY


def test_holiday_beats_saturday():
    assert HOLIDAY_ON_SATURDAY.weekday() == 5
    assert day_zones_config(_data(True), WEEKDAY, HOLIDAY_ON_SATURDAY) == SUNDAY
    assert day_zones_config(_data(False), WEEKDAY, HOLIDAY_ON_SATURDAY) == SATURDAY


def test_holiday_without_sunday_override_uses_base():
    data = {CONF_TARIFF_ZONES: WEEKDAY, CONF_HOLIDAYS: "national"}
    assert day_zones_config(data, WEEKDAY, HOLIDAY_ON_WEEKDAY) == WEEKDAY


def test_state_specific_holiday_is_not_national():
    assert not is_national_holiday(date(2026, 6, 4))  # Fronleichnam


def test_distinct_zone_names_orders_by_first_appearance():
    assert distinct_zone_names(WEEKDAY, SATURDAY, SUNDAY) == ["NT", "HT"]
    assert distinct_zone_names(SUNDAY, WEEKDAY) == ["NT", "HT"]


# --- form validation ---------------------------------------------------------


def test_optional_zones_accept_empty_and_single_entry():
    assert _parse_optional_tariff_zones([]) == []
    assert _parse_optional_tariff_zones(["00:00 NT"]) == [
        {"time": "00:00", "name": "NT"}
    ]


def test_optional_zones_still_validate_format():
    with pytest.raises(ZoneDefinitionError):
        _parse_optional_tariff_zones(["06:00 NT"])  # must start at 00:00


# --- sensors -------------------------------------------------------------------


def test_sensors_cover_zone_names_and_switches_of_all_layouts():
    descs = _dynamic_descriptions(WEEKDAY, SATURDAY, SUNDAY)
    slots = [d for d in descs if d.key.startswith("daily_consumption_slot_")]
    switches = [d for d in descs if d.key.startswith("meter_consumption_switch_")]
    assert [d.translation_placeholders["zone_name"] for d in slots] == ["NT", "HT"]
    assert len(switches) == 2  # longest layout has 3 entries -> 2 inner switches


def test_zone_missing_on_the_day_is_zero_not_unknown():
    """Sunday is all NT: HT consumed 0 kWh, the HT sensor must not be unknown."""
    dd = DailyData(
        date=SUNDAY_D,
        import_boundaries=[1000.0, 1004.0],
        export_midnight=0.0,
        export_next_midnight=0.0,
        zone_totals={"NT": 4.0},
        daily_import_total=4.0,
        daily_export_total=0.0,
    )
    data = SmgwCoordinator._daily_data_to_dict(dd, {"NT": 1, "HT": 2})
    assert data["daily_consumption_slot_1"] == 4.0
    assert data["daily_consumption_slot_2"] == 0.0


def test_daily_data_to_dict_without_names_keeps_old_behaviour():
    dd = DailyData(
        date=MONDAY_D,
        import_boundaries=[1000.0, 1001.0, 1010.0],
        export_midnight=0.0,
        export_next_midnight=0.0,
        zone_totals={"NT": 1.0, "HT": 9.0},
        daily_import_total=10.0,
        daily_export_total=0.0,
    )
    data = SmgwCoordinator._daily_data_to_dict(dd)
    assert data["daily_consumption_slot_2"] == 9.0


# --- coordinator ---------------------------------------------------------------


async def _coordinator(hass: HomeAssistant, stub, **data) -> SmgwCoordinator:
    await hass.config.async_set_time_zone("Europe/Berlin")
    entry = MockConfigEntry(
        domain=DOMAIN, data={CONF_METER_ID: "M", **_data(), **data}
    )
    entry.add_to_hass(hass)
    return SmgwCoordinator(hass, entry, stub)


def _dd(day: date, totals: dict[str, float]) -> DailyData:
    total = sum(totals.values())
    return DailyData(
        date=day,
        import_boundaries=[1000.0, 1000.0 + total],
        export_midnight=0.0,
        export_next_midnight=0.0,
        zone_totals=totals,
        daily_import_total=total,
        daily_export_total=0.0,
    )


class _FetchStub:
    """Client stub for the daily fetch; remembers the zones it was given."""

    zones = None

    def __init__(self, *, result=None, exc=None):
        self._result = result
        self._exc = exc
        self.calls = 0

    async def close(self):
        pass


    async def async_fetch_daily_data(self, target_date, zones=None, target_meter_id=None):
        self.calls += 1
        self.zones = zones
        if self._exc is not None:
            raise self._exc
        return self._result


@pytest.mark.parametrize(
    ("now", "expected"),
    [
        ("2026-10-06 00:15:00", WEEKDAY),  # fetches Monday 10-05
        ("2026-10-11 00:15:00", SATURDAY),  # fetches Saturday 10-10
        ("2026-10-12 00:15:00", SUNDAY),  # fetches Sunday 10-11
        ("2026-10-04 00:15:00", SUNDAY),  # fetches 10-03: Saturday AND holiday
    ],
)
async def test_fetch_splits_each_day_by_its_own_layout(hass, now, expected):
    stub = _FetchStub(result=_dd(date(2026, 1, 1), {"NT": 1.0}))
    coord = await _coordinator(hass, stub)
    with freeze_time(now):
        await coord._async_do_daily_fetch()
    assert [(t.strftime("%H:%M"), n) for t, n in stub.zones] == [
        (z["time"], z["name"]) for z in expected
    ]


async def test_sunday_fetch_publishes_ht_as_zero(hass: HomeAssistant):
    stub = _FetchStub(result=_dd(date(2026, 10, 11), {"NT": 3.0}))
    coord = await _coordinator(hass, stub)
    with freeze_time("2026-10-12 00:15:00"):
        await coord._async_do_daily_fetch()
    assert coord.data["daily_consumption_slot_1"] == 3.0
    assert coord.data["daily_consumption_slot_2"] == 0.0


async def test_zone_resolver_combines_schedule_and_day_types(hass: HomeAssistant):
    """Dated Mon-Fri change from 2027 on; Saturday override stays; holiday = Sunday."""
    new_week = [
        {"time": "00:00", "name": "NT"},
        {"time": "07:00", "name": "HT"},
        {"time": "21:00", "name": "NT"},
    ]
    schedule = [
        {"valid_from": None, "zones": WEEKDAY},
        {"valid_from": "2027-01-01", "zones": new_week},
    ]
    coord = await _coordinator(
        hass, _FetchStub(result=None), **{CONF_ZONE_SCHEDULE: schedule}
    )
    resolve = coord.zone_resolver()

    def layout(day: date) -> list[tuple[str, str]]:
        return [(t.strftime("%H:%M"), n) for t, n in resolve(day)]

    as_pairs = lambda zones: [(z["time"], z["name"]) for z in zones]  # noqa: E731
    assert layout(date(2026, 12, 30)) == as_pairs(WEEKDAY)  # Wed, old layout
    assert layout(date(2027, 1, 5)) == as_pairs(new_week)  # Tue, new layout
    assert layout(date(2027, 1, 9)) == as_pairs(SATURDAY)  # Sat override
    assert layout(date(2027, 1, 6)) == as_pairs(new_week)  # Epiphany: state holiday only
    assert layout(date(2027, 5, 3)) == as_pairs(new_week)  # Monday
    assert layout(date(2027, 5, 1)) == as_pairs(SUNDAY)  # Labour Day (Saturday)


async def test_slot_names_cover_all_day_types(hass: HomeAssistant):
    coord = await _coordinator(
        hass,
        _FetchStub(result=None),
        **{CONF_TARIFF_ZONES_SUNDAY: [{"time": "00:00", "name": "SPITZE"}]},
    )
    assert coord._slot_zone_names == {"NT": 1, "HT": 2, "SPITZE": 101}


async def test_fetch_remembers_the_sensor_set(hass: HomeAssistant):
    stub = _FetchStub(result=_dd(date(2026, 10, 11), {"NT": 3.0}))
    coord = await _coordinator(hass, stub)
    with freeze_time("2026-10-12 00:15:00"):
        await coord._async_do_daily_fetch()
    assert coord.data["_slot_zone_names"] == {"NT": 1, "HT": 2}


async def test_day_types_listed_for_export(hass: HomeAssistant):
    coord = await _coordinator(hass, _FetchStub(result=None))
    labels = [
        label
        for label, _ in coord.zone_day_types(date(2026, 10, 1), date(2026, 10, 31))
    ]
    assert labels == ["Samstag", "Sonntag", "Feiertage bundeseinheitlich"]
    plain = await _coordinator(
        hass,
        _FetchStub(result=None),
        **{CONF_TARIFF_ZONES_SATURDAY: [], CONF_TARIFF_ZONES_SUNDAY: [],
           CONF_HOLIDAYS: "off"},
    )
    assert plain.zone_day_types(date(2026, 10, 1), date(2026, 10, 31)) == []


# --- aggregation + export ------------------------------------------------------


def _readings() -> list[MeterReading]:
    """Sat 2026-10-10 (HT 06-13) and Sun 2026-10-11 (all NT)."""
    def imp(d, h, m, v):
        return MeterReading(datetime(2026, 10, d, h, m, 1), OBIS_IMPORT, v, "kWh", "valid")

    return [
        imp(10, 0, 0, 1000.0),
        imp(10, 6, 0, 1001.0),
        imp(10, 13, 0, 1004.0),
        imp(11, 0, 0, 1005.0),
        imp(12, 0, 0, 1007.0),
    ]


def _resolver(day: date):
    data = _data()
    base = [{"time": z["time"], "name": z["name"]} for z in WEEKDAY]
    return [
        (time.fromisoformat(z["time"]), z["name"])
        for z in day_zones_config(data, base, day)
    ]


def test_aggregation_splits_weekend_days_by_their_layout():
    saturday, sunday = build_daily_summary(_readings(), _resolver)
    assert saturday.zone_consumptions == {"NT": 2.0, "HT": 3.0}
    assert sunday.zone_consumptions == {"NT": 2.0}
    assert sunday.consumption_total == 2.0


def test_xlsx_has_union_columns_and_documents_day_types(tmp_path):
    summary = build_daily_summary(_readings(), _resolver)
    meta = {
        "meter_id": "M",
        "from": "2026-10-10",
        "to": "2026-10-12",
        "zones": [(z["time"], z["name"]) for z in WEEKDAY],
        "day_type_zones": [
            {"label": "Samstag", "zones": [(z["time"], z["name"]) for z in SATURDAY]},
            {"label": "Sonntag", "zones": [(z["time"], z["name"]) for z in SUNDAY]},
        ],
    }
    path = tmp_path / "x.xlsx"
    write_xlsx(path, _readings(), summary, meta)
    wb = load_workbook(path)
    headers = [c.value for c in wb["Tarifzonen"][1]]
    # 06:00, 13:00 and 22:00 all occur in some layout -> three Bezug columns
    for t in ("06:00", "13:00", "22:00"):
        assert f"Bezug {t} (kWh)" in headers
    info = "\n".join(
        str(row[0].value) for row in wb["Definition"].iter_rows() if row[0].value
    )
    assert "Samstag (abweichend)" in info
    assert "Sonntag (abweichend)" in info


def _tarifzonen_rows(path) -> dict[str, dict]:
    """The "Tarifzonen" sheet as {date: {column header: value}}."""
    sheet = load_workbook(path)["Tarifzonen"]
    headers = [c.value for c in sheet[1]]
    return {
        row[0].value: dict(zip(headers, (c.value for c in row)))
        for row in sheet.iter_rows(min_row=2)
    }


def test_xlsx_zone_missing_on_a_day_is_zero_like_its_sensor(tmp_path):
    summary = build_daily_summary(_readings(), _resolver)
    path = tmp_path / "x.xlsx"
    write_xlsx(path, _readings(), summary, {"zones": [("00:00", "NT"), ("06:00", "HT")]})
    sunday = _tarifzonen_rows(path)["2026-10-11"]
    assert sunday["Verbrauch NT (kWh)"] == 2.0
    assert sunday["Verbrauch HT (kWh)"] == 0


def test_xlsx_missing_reading_stays_an_empty_cell(tmp_path):
    readings = [
        r for r in _readings() if r.timestamp != datetime(2026, 10, 10, 13, 0, 1)
    ]
    summary = build_daily_summary(readings, _resolver)
    path = tmp_path / "x.xlsx"
    write_xlsx(path, readings, summary, {"zones": [("00:00", "NT"), ("06:00", "HT")]})
    saturday = _tarifzonen_rows(path)["2026-10-10"]
    assert saturday["Verbrauch HT (kWh)"] is None
    assert saturday["Verbrauch NT (kWh)"] is None


def test_xlsx_names_the_holidays_within_the_exported_days(tmp_path):
    summary = build_daily_summary(_readings(), _resolver)
    path = tmp_path / "x.xlsx"
    write_xlsx(
        path, _readings(), summary,
        {
            "zones": [("00:00", "NT"), ("06:00", "HT")],
            "holiday_days": ["2026-10-11", "2026-10-12"],
        },
    )
    info = [
        str(row[0].value)
        for row in load_workbook(path)["Definition"].iter_rows()
        if row[0].value
    ]
    # The range may end on a holiday without a complete day in the workbook.
    assert "Als Feiertag behandelte Tage: 2026-10-11" in info


async def test_holidays_in_range_follow_each_period(hass: HomeAssistant):
    """Off until the end of 2026, nationwide from 2027 on."""
    schedule = [
        {"valid_from": None, "zones": WEEKDAY, "saturday": [], "sunday": SUNDAY,
         "holidays": "off"},
        {"valid_from": "2027-01-01", "zones": WEEKDAY, "saturday": [],
         "sunday": SUNDAY, "holidays": "national"},
    ]
    coord = await _coordinator(
        hass, _FetchStub(result=None), **{CONF_ZONE_SCHEDULE: schedule}
    )
    assert coord.holidays_in_range(date(2026, 12, 20), date(2027, 1, 10)) == [
        date(2027, 1, 1)
    ]


# --- restart with a cached day -----------------------------------------------


def _store(hass_storage: dict, entry: MockConfigEntry, data: dict) -> None:
    key = f"{DOMAIN}_{entry.entry_id}"
    hass_storage[key] = {"version": STORE_VERSION, "minor_version": 1, "key": key, "data": data}


def _cached(day: str, zones: list[dict], **extra) -> dict:
    return {
        "date": day,
        "daily_consumption_total": 3.0,
        "daily_feedin_total": 0.0,
        "meter_consumption_prev_day_close": 1003.0,
        "meter_feedin_prev_day_close": 0.0,
        "daily_consumption_slot_1": 3.0,
        "daily_consumption_slot_2": 0.0,
        "_tariff_zones": zones,
        **extra,
    }


async def _restart(
    hass, hass_storage, cached: dict, now: str, stub=None, **data
) -> SmgwCoordinator:
    await hass.config.async_set_time_zone("Europe/Berlin")
    entry = MockConfigEntry(
        domain=DOMAIN, data={CONF_METER_ID: "M", **_data(), **data}
    )
    entry.add_to_hass(hass)
    _store(hass_storage, entry, cached)
    # A gateway that must not be asked unless a test says otherwise.
    stub = stub or _FetchStub(exc=AssertionError("unexpected fetch"))
    coord = SmgwCoordinator(hass, entry, stub)
    with freeze_time(now):
        await coord.async_setup()
    return coord


async def test_restart_keeps_sunday_values_published(hass, hass_storage):
    """Sunday's cache must survive a restart on Monday (layouts unchanged)."""
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-11", SUNDAY, _slot_zone_names={"NT": 1, "HT": 2}),
        "2026-10-12 09:00:00",
    )
    assert coord.data["daily_consumption_slot_1"] == 3.0
    assert coord.data["daily_consumption_slot_2"] == 0.0
    await coord.async_unload()


async def test_restart_with_legacy_cache_keeps_values(hass, hass_storage):
    """A cache written before day types existed (no _slot_zone_names)."""
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-05", WEEKDAY),  # Monday, computed with Mon-Fri zones
        "2026-10-06 09:00:00",
    )
    assert coord.data["daily_consumption_slot_1"] == 3.0
    await coord.async_unload()


async def test_changed_sensor_set_withholds_cached_zone_values(
    hass, hass_storage
):
    """If the sensor set (zone names) differs from the cache, values go."""
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-11", SUNDAY, _slot_zone_names={"NT": 1, "HT": 2, "SPITZE": 101}),
        "2026-10-12 09:00:00",
        stub=_FetchStub(exc=SmgwConnectionError("down")),
    )
    assert "daily_consumption_slot_1" not in coord.data
    assert coord.data["daily_consumption_total"] == 3.0
    await coord.async_unload()


async def test_new_weekend_only_name_refetches_to_fill_the_sensors(
    hass, hass_storage
):
    """The cached Monday is still valid, but the new sensor needs its value."""
    stub = _FetchStub(result=_dd(date(2026, 10, 5), {"NT": 2.0, "HT": 1.0}))
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-05", WEEKDAY, _slot_zone_names={"NT": 1, "HT": 2}),
        "2026-10-06 09:00:00",
        stub=stub,
        **{CONF_TARIFF_ZONES_SUNDAY: [{"time": "00:00", "name": "SPITZE"}]},
    )
    assert stub.calls == 1
    assert coord.data["daily_consumption_slot_1"] == 2.0
    assert coord.data["daily_consumption_slot_2"] == 1.0
    assert coord.data["daily_consumption_slot_101"] == 0.0
    await coord.async_unload()


async def test_legacy_sunday_cache_is_withheld_after_adding_sunday_zones(
    hass, hass_storage
):
    """An old cache split that Sunday by the zones valid every day back then."""
    stub = _FetchStub(exc=SmgwConnectionError("down"))
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-11", WEEKDAY),
        "2026-10-12 09:00:00",
        stub=stub,
    )
    assert stub.calls == 1
    assert "daily_consumption_slot_1" not in coord.data
    assert coord.data["daily_consumption_total"] == 3.0
    await coord.async_unload()


async def test_legacy_holiday_cache_is_withheld_when_holidays_count(
    hass, hass_storage
):
    """25 Dec is a Friday: split as a weekday by the old version."""
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-12-25", WEEKDAY),
        "2026-12-26 09:00:00",
        stub=_FetchStub(exc=SmgwConnectionError("down")),
    )
    assert "daily_consumption_slot_1" not in coord.data
    await coord.async_unload()


async def test_changed_times_withhold_the_day_they_affect(hass, hass_storage):
    """Same zone names, other times: the old split must not stay visible."""
    other_times = [
        {"time": "00:00", "name": "NT"},
        {"time": "07:00", "name": "HT"},
        {"time": "22:00", "name": "NT"},
    ]
    stub = _FetchStub(exc=SmgwConnectionError("down"))
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-05", WEEKDAY, _slot_zone_names={"NT": 1, "HT": 2}),
        "2026-10-06 09:00:00",
        stub=stub,
        **{CONF_TARIFF_ZONES: other_times},
    )
    assert stub.calls == 1
    assert "daily_consumption_slot_1" not in coord.data
    await coord.async_unload()


async def test_changed_sunday_times_leave_a_weekday_cache_alone(
    hass, hass_storage
):
    """Nothing about the cached Monday changed: no withholding, no fetch."""
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-05", WEEKDAY, _slot_zone_names={"NT": 1, "HT": 2}),
        "2026-10-06 09:00:00",
        **{
            CONF_TARIFF_ZONES_SUNDAY: [
                {"time": "00:00", "name": "NT"},
                {"time": "12:00", "name": "HT"},
            ]
        },
    )
    assert coord.data["daily_consumption_slot_1"] == 3.0
    await coord.async_unload()


async def test_old_day_stays_published_after_a_scheduled_time_change(
    hass, hass_storage
):
    """31 Dec is correctly split by the old times, and the names did not move."""
    new_week = [
        {"time": "00:00", "name": "NT"},
        {"time": "07:00", "name": "HT"},
        {"time": "21:00", "name": "NT"},
    ]
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-12-31", WEEKDAY, _slot_zone_names={"NT": 1, "HT": 2}),
        "2027-01-01 09:00:00",
        **{
            CONF_TARIFF_ZONES: new_week,
            CONF_ZONE_SCHEDULE: [
                {"valid_from": None, "zones": WEEKDAY},
                {"valid_from": "2027-01-01", "zones": new_week},
            ],
        },
    )
    assert coord.data["daily_consumption_slot_1"] == 3.0
    await coord.async_unload()


async def test_withheld_values_do_not_count_as_already_fetched(
    hass, hass_storage
):
    """A later fetch of the same day must not be skipped."""
    stub = _FetchStub(exc=SmgwConnectionError("down"))
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-11", WEEKDAY),
        "2026-10-12 09:00:00",
        stub=stub,
    )
    stub._exc = None
    stub._result = _dd(date(2026, 10, 11), {"NT": 3.0})
    with freeze_time("2026-10-12 09:30:00"):
        assert await coord._async_do_daily_fetch() is True
    assert coord.data["daily_consumption_slot_1"] == 3.0
    assert coord.data["daily_consumption_slot_2"] == 0.0
    await coord.async_unload()


async def test_changed_holiday_mode_keeps_cached_zone_values(
    hass, hass_storage
):
    """Same zone names, a day the setting does not touch: values stay."""
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-11", SUNDAY, _slot_zone_names={"NT": 1, "HT": 2}),
        "2026-10-12 09:00:00",
        **{CONF_HOLIDAYS: "by"},
    )
    assert coord.data["daily_consumption_slot_1"] == 3.0
    await coord.async_unload()


# --- dated rules and holiday modes ---------------------------------------------


def test_holiday_modes_off_national_state():
    epiphany = date(2027, 1, 6)  # Wednesday, only some states
    corpus = date(2026, 6, 4)  # Thursday, only some states
    assert not day_rules.is_holiday("off", date(2026, 12, 25))
    assert day_rules.is_holiday("national", date(2026, 12, 25))
    assert not day_rules.is_holiday("national", epiphany)
    assert day_rules.is_holiday("BY", epiphany)
    assert day_rules.is_holiday("BY", corpus)
    assert not day_rules.is_holiday("NW", epiphany)
    assert day_rules.is_holiday("NW", corpus)


async def test_dated_rules_replace_weekend_from_their_date(hass: HomeAssistant):
    """From 2027-01-01 Saturday is dropped and Bavarian holidays apply."""
    new_week = [{"time": "00:00", "name": "NT"}, {"time": "07:00", "name": "HT"}]
    schedule = [
        {"valid_from": None, "zones": WEEKDAY},
        {
            "valid_from": "2027-01-01",
            "zones": new_week,
            "saturday": [],
            "sunday": SUNDAY,
            "holidays": "BY",
        },
    ]
    coord = await _coordinator(
        hass, _FetchStub(result=None), **{CONF_ZONE_SCHEDULE: schedule}
    )
    resolve = coord.zone_resolver()
    await coord.async_preload_holidays()

    def layout(day: date) -> list[tuple[str, str]]:
        return [(t.strftime("%H:%M"), n) for t, n in resolve(day)]

    as_pairs = lambda zones: [(z["time"], z["name"]) for z in zones]  # noqa: E731
    assert layout(date(2026, 12, 19)) == as_pairs(SATURDAY)  # old rules
    assert layout(date(2027, 1, 2)) == as_pairs(new_week)  # Saturday dropped
    assert layout(date(2027, 1, 3)) == as_pairs(SUNDAY)
    assert layout(date(2027, 1, 6)) == as_pairs(SUNDAY)  # Epiphany, Bavaria


def test_schedule_with_change_carries_rules():
    from custom_components.smgw_han import zone_schedule

    sched = zone_schedule.with_change(
        None,
        WEEKDAY,
        date(2027, 1, 1),
        WEEKDAY,
        current_rules={"saturday": SATURDAY, "sunday": SUNDAY, "holidays": "off"},
        rules={"saturday": [], "sunday": [], "holidays": "national"},
    )
    assert sched[0]["saturday"] == SATURDAY
    assert sched[1]["saturday"] == []
    assert sched[1]["holidays"] == "national"


# --- regressions found in review -----------------------------------------------


def test_with_change_writes_inherited_rules_into_older_entries():
    """An old-style schedule must not inherit the NEW rules into its history."""
    from custom_components.smgw_han import zone_schedule

    old = [
        {"valid_from": None, "zones": WEEKDAY},
        {"valid_from": "2026-01-01", "zones": WEEKDAY},
    ]
    sched = zone_schedule.with_change(
        old,
        WEEKDAY,
        date(2027, 1, 1),
        WEEKDAY,
        current_rules={"saturday": SATURDAY, "sunday": SUNDAY, "holidays": "national"},
        rules={"saturday": [], "sunday": [], "holidays": "off"},
    )
    layout = zone_schedule.layout_for_day
    fallback = {"zones": WEEKDAY, "saturday": [], "sunday": [], "holidays": "off"}
    assert layout(sched, fallback, date(2026, 6, 4))["holidays"] == "national"
    assert layout(sched, fallback, date(2026, 6, 6))["saturday"] == SATURDAY
    assert layout(sched, fallback, date(2027, 6, 5))["saturday"] == []


def test_selector_keys_are_hassfest_safe():
    import json
    import re
    from pathlib import Path

    import custom_components.smgw_han as pkg

    base = Path(pkg.__file__).parent
    for name in ("strings.json", "translations/en.json", "translations/de.json"):
        options = json.loads((base / name).read_text(encoding="utf-8"))[
            "selector"
        ]["holiday_mode"]["options"]
        assert all(re.fullmatch(r"[a-z0-9-_]+", key) for key in options)
    from custom_components.smgw_han.const import HOLIDAY_STATES

    assert set(HOLIDAY_STATES) | {"off", "national"} == set(options)


def test_bavaria_includes_assumption_day():
    assert day_rules.is_holiday("by", date(2026, 8, 15))
    assert not day_rules.is_holiday("nw", date(2026, 8, 15))
    assert day_rules.holiday_mode({CONF_HOLIDAYS: "BY"}) == "by"


# --- schedule rules and scheduled renames -------------------------------------

RULES_OLD = {"saturday": [], "sunday": [], "holidays": "off"}
RULES_NEW = {
    "saturday": [{"time": "00:00", "name": "NT"}],
    "sunday": [{"time": "00:00", "name": "NT"}],
    "holidays": "by",
}
ZONES_WD = [{"time": "00:00", "name": "NT"}, {"time": "06:00", "name": "HT"}]


def test_with_change_later_ruleless_entry_inherits_the_new_rules():
    """A pending zone-only entry must not be pinned to the replaced rules."""
    schedule = [
        {"valid_from": None, "zones": ZONES_WD},
        {"valid_from": "2027-01-01", "zones": ZONES_WD},  # no rule keys
    ]
    result = zone_schedule.with_change(
        schedule, ZONES_WD, date(2026, 10, 26), ZONES_WD,
        current_rules=RULES_OLD, rules=RULES_NEW,
    )
    by_date = {item["valid_from"]: item for item in result}
    assert by_date[None]["holidays"] == "off"          # history keeps the old rules
    assert by_date["2026-10-26"]["holidays"] == "by"
    assert by_date["2027-01-01"]["holidays"] == "by"   # inherits the NEW rules
    assert by_date["2027-01-01"]["sunday"] == RULES_NEW["sunday"]


def test_with_change_planner_keeps_later_entry_on_new_rules():
    schedule = [
        {"valid_from": None, "zones": ZONES_WD},
        {"valid_from": "2027-06-01", "zones": ZONES_WD},
    ]
    result = zone_schedule.with_change(
        schedule, ZONES_WD, date(2027, 1, 1), ZONES_WD,
        current_rules=RULES_OLD, rules=RULES_NEW,
    )
    by_date = {item["valid_from"]: item for item in result}
    assert by_date[None]["sunday"] == []
    assert by_date["2027-06-01"]["sunday"] == RULES_NEW["sunday"]
    assert by_date["2027-06-01"]["holidays"] == "by"


def test_schedule_holiday_codes_are_normalized():
    result = zone_schedule.normalize(
        [{"valid_from": None, "zones": ZONES_WD, "holidays": " BY "}]
    )
    assert result[0]["holidays"] == "by"


def test_fetch_after_scheduled_rename_feeds_sensors_by_number():
    """Last day of the OLD layout, whose names are not today's slot names."""
    dd = DailyData(
        date=date(2026, 12, 31),
        daily_import_total=10.0,
        daily_export_total=0.0,
        zone_totals={"Go": 4.0, "Standard": 6.0},
        import_boundaries=[1.0, 5.0, 11.0],
        export_midnight=0.0,
        export_next_midnight=0.0,
    )
    data = SmgwCoordinator._daily_data_to_dict(
        dd, {"NT": 1, "HT": 2}, day_slots={"Go": 1, "Standard": 2}
    )
    assert data["daily_consumption_slot_1"] == 4.0
    assert data["daily_consumption_slot_2"] == 6.0


def test_fetch_with_same_names_still_fills_by_name():
    dd = DailyData(
        date=date(2026, 10, 11),
        daily_import_total=3.0,
        daily_export_total=0.0,
        zone_totals={"NT": 3.0},
        import_boundaries=[1.0, 4.0],
        export_midnight=0.0,
        export_next_midnight=0.0,
    )
    data = SmgwCoordinator._daily_data_to_dict(
        dd, {"HT": 1, "NT": 2}, day_slots={"NT": 1}
    )
    assert data["daily_consumption_slot_1"] == 0.0
    assert data["daily_consumption_slot_2"] == 3.0


WEEKEND_ONLY = [
    {"time": "00:00", "name": "NT"},
    {"time": "08:00", "name": "Weekend"},
]


def test_weekend_only_names_get_their_own_number_range():
    assert day_rules.zone_slots(WEEKDAY, SATURDAY, SUNDAY) == {"NT": 1, "HT": 2}
    assert day_rules.zone_slots(WEEKDAY, WEEKEND_ONLY, SUNDAY) == {
        "NT": 1,
        "HT": 2,
        "Weekend": 101,
    }


def _slot_keys(*layouts) -> dict[str, str]:
    return {
        d.translation_placeholders["zone_name"]: d.key
        for d in _dynamic_descriptions(*layouts)
        if d.key.startswith("daily_consumption_slot_")
    }


def test_new_mon_fri_name_does_not_move_a_weekend_only_sensor():
    before = _slot_keys(WEEKDAY, WEEKEND_ONLY, WEEKEND_ONLY)
    with_peak = [*WEEKDAY, {"time": "23:00", "name": "Peak"}]
    after = _slot_keys(with_peak, WEEKEND_ONLY, WEEKEND_ONLY)
    assert before["Weekend"] == after["Weekend"] == "daily_consumption_slot_101"
    assert after["Peak"] == "daily_consumption_slot_3"


def test_scheduled_rename_keeps_a_weekend_only_value_on_its_sensor():
    """Rename plus new Mon-Fri name: the old Sunday must not feed the new name."""
    dd = DailyData(
        date=date(2027, 1, 3),
        daily_import_total=13.0,
        daily_export_total=0.0,
        zone_totals={"HT": 8.0, "Weekend": 5.0},
        import_boundaries=[1.0, 9.0, 14.0],
        export_midnight=0.0,
        export_next_midnight=0.0,
    )
    data = SmgwCoordinator._daily_data_to_dict(
        dd,
        {"NT": 1, "Tag": 2, "Peak": 3, "Weekend": 101},
        day_slots={"NT": 1, "HT": 2, "Weekend": 101},
    )
    assert data["daily_consumption_slot_1"] == 0.0
    assert data["daily_consumption_slot_2"] == 8.0
    assert data["daily_consumption_slot_3"] == 0.0
    assert data["daily_consumption_slot_101"] == 5.0


async def test_cache_of_the_first_beta_is_refetched_once(hass, hass_storage):
    """v3.8.0-beta.1 stored its slot names as a list without numbers."""
    stub = _FetchStub(result=_dd(date(2026, 10, 5), {"NT": 2.0, "HT": 1.0}))
    coord = await _restart(
        hass, hass_storage,
        _cached("2026-10-05", WEEKDAY, _slot_zone_names=["NT", "HT"]),
        "2026-10-06 09:00:00",
        stub=stub,
    )
    assert stub.calls == 1
    assert coord.data["_slot_zone_names"] == {"NT": 1, "HT": 2}
    await coord.async_unload()


def test_optional_zone_field_ignores_cleared_chips():
    from custom_components.smgw_han.config_flow import (
        _parse_optional_tariff_zones,
    )

    assert _parse_optional_tariff_zones([""]) == []
    assert _parse_optional_tariff_zones(["", "00:00 NT"]) == [
        {"time": "00:00", "name": "NT"}
    ]
